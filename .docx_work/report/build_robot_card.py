from copy import deepcopy
from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.shared import Inches
from PIL import Image


REFERENCE = Path(r"C:\Users\Sanchive Kumar\Downloads\Contractsentry one page report.docx")
OUTPUT = Path(r"C:\Users\Sanchive Kumar\robot\AI Museum Welcome Robot Showcase Card.docx")
SOURCE_IMAGE = Path(r"C:\Users\Sanchive Kumar\robot\.docx_work\report\welcome_robot_illustration.png")
CARD_IMAGE = Path(r"C:\Users\Sanchive Kumar\robot\.docx_work\report\welcome_robot_card_image.png")


def copy_run_style(source, destination):
    if source is not None and source._r.rPr is not None:
        destination._r.get_or_add_rPr().append(deepcopy(source._r.rPr))


def replace_plain(paragraph, text):
    source = paragraph.runs[0] if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(text)
    copy_run_style(source, run)


def replace_labeled(paragraph, label, value):
    label_source = paragraph.runs[0] if paragraph.runs else None
    value_source = paragraph.runs[1] if len(paragraph.runs) > 1 else label_source
    paragraph.clear()
    label_run = paragraph.add_run(label)
    copy_run_style(label_source, label_run)
    value_run = paragraph.add_run(value)
    copy_run_style(value_source, value_run)


def replace_bullet(paragraph, text):
    bullet_source = paragraph.runs[0] if paragraph.runs else None
    text_source = paragraph.runs[1] if len(paragraph.runs) > 1 else bullet_source
    paragraph.clear()
    bullet_run = paragraph.add_run("•  ")
    copy_run_style(bullet_source, bullet_run)
    text_run = paragraph.add_run(text)
    copy_run_style(text_source, text_run)


def prepare_card_image():
    source = Image.open(SOURCE_IMAGE).convert("RGB")
    canvas = Image.new("RGB", (1200, 1000), "white")
    source.thumbnail((666, 1000), Image.Resampling.LANCZOS)
    x = (canvas.width - source.width) // 2
    y = (canvas.height - source.height) // 2
    canvas.paste(source, (x, y))
    canvas.save(CARD_IMAGE, quality=95)


def replace_picture(paragraph, image, width, height):
    paragraph.clear()
    paragraph.add_run().add_picture(str(image), width=width, height=height)


def main():
    copyfile(REFERENCE, OUTPUT)
    prepare_card_image()
    document = Document(OUTPUT)

    title_cell = document.tables[2].cell(0, 0)
    replace_plain(title_cell.paragraphs[0], "PROJECT NAME: AI MUSEUM WELCOME ROBOT - VISION AND VOICE GUIDE")
    replace_labeled(title_cell.paragraphs[1], "Domain: ", "Embodied AI / Computer Vision / Robotics")

    narrative = document.tables[3].cell(0, 1).paragraphs
    replace_plain(
        narrative[1],
        "Museum visitors need quick, friendly guidance, but a human guide cannot always greet everyone, answer repeated event questions and manage safe movement.",
    )
    replace_bullet(narrative[3], "Core AI: YOLO26n vision, Sarvam voice AI and ESP32 safety control")
    replace_bullet(
        narrative[4],
        "Working: Camera -> Visitor / VIP detection -> Safe greeting -> Voice question -> Event-grounded answer",
    )
    replace_bullet(
        narrative[5],
        "Data: Event schedule, facilities and 17-project catalog; consented VIP profiles stored locally",
    )
    replace_bullet(
        narrative[8],
        "Vision-guided greeting with YOLO26n and three-frame VIP verification",
    )
    replace_bullet(
        narrative[9],
        "Speech-to-text, event-grounded Q&A and natural TTS responses",
    )
    replace_bullet(
        narrative[10],
        "Safety-first movement: ultrasonic stop, wheel encoders, E-stop and watchdog",
    )

    main_image_slot = document.tables[3].cell(0, 0).tables[0].cell(0, 0).paragraphs[0]
    replace_picture(main_image_slot, CARD_IMAGE, Inches(2.84), Inches(2.36))

    qr_slot = document.tables[3].cell(0, 0).tables[1].cell(0, 0).tables[0].cell(0, 0).paragraphs[0]
    replace_plain(qr_slot, "LIVE\nDEMO")
    status_paragraphs = document.tables[3].cell(0, 0).tables[1].cell(0, 1).paragraphs
    replace_plain(status_paragraphs[0], "Status:\n ☑ Prototype")
    for extra_paragraph in status_paragraphs[1:]:
        replace_plain(extra_paragraph, "")

    technologies = document.tables[4].cell(0, 0).tables[0].rows[0].cells
    for cell, label in zip(
        technologies,
        ["Python", "FastAPI", "ESP32", "YOLO26n", "OpenCV", "Sarvam AI", "DroidCam"],
    ):
        replace_plain(cell.paragraphs[0], label)

    applications = document.tables[5].cell(0, 0).paragraphs
    replace_bullet(applications[1], "Museum reception and visitor guidance")
    replace_bullet(applications[2], "Event information and exhibit Q&A")
    replace_bullet(applications[3], "Campus, lab and exhibition assistance")

    document.save(OUTPUT)


if __name__ == "__main__":
    main()
